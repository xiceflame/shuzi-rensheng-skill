#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
附件抽取器：把 focus-media / raw/docs 里的 PDF/Office/表格/文本 **抽成纯文本**，
供 LLM 读取与整理（不再只有聊天文字）。

- 输入：一个目录（如 ~/wechat-export/focus-media）或单个文件
- 输出：~/wechat-export/media-text/<相对路径>.txt  +  _manifest.tsv
- 增量：目标已存在且比源新 → 跳过
- 无法抽文本的（图片/音视频/zip 等）在 manifest 里标 agent-readable（由 agent 用 pdf/view_image/转写工具处理）

用法：
    媒体venv/bin/python media_extract.py ~/wechat-export/focus-media
    媒体venv/bin/python media_extract.py <单文件>
"""
import os
import sys
import subprocess
import datetime

SRC_ROOT = os.path.expanduser("~/wechat-export")
OUT_ROOT = os.path.join(SRC_ROOT, "media-text")
TEXT_EXT = {".txt", ".md", ".csv", ".json", ".tsv", ".log", ".ino", ".py", ".js", ".sh", ".yaml", ".yml", ".xml", ".html", ".tex"}
IMG_EXT = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".heic", ".hevc", ".bmp", ".tif", ".tiff"}
AV_EXT = {".m4a", ".mp3", ".wav", ".aac", ".amr", ".mp4", ".mov", ".mkv", ".avi"}
ARC_EXT = {".zip", ".rar", ".7z", ".tar", ".gz"}
UNSUPPORTED = {".ppt", ".pages", ".numbers", ".key", ".blend", ".stl", ".obj", ".dwg"}


def read_text(p, limit=None):
    with open(p, encoding="utf-8", errors="replace") as f:
        t = f.read()
    return t[:limit] if limit else t


def extract_pdf(p):
    from pypdf import PdfReader
    r = PdfReader(p)
    out = []
    for i, pg in enumerate(r.pages):
        try:
            out.append("--- p%d ---\n%s" % (i + 1, pg.extract_text() or ""))
        except Exception as e:
            out.append("--- p%d ---（提取失败：%s）" % (i + 1, e))
    txt = "\n".join(out)
    return txt, len(r.pages)


def extract_docx(p):
    import docx
    d = docx.Document(p)
    parts = [para.text for para in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts), None


def extract_pptx(p):
    from pptx import Presentation
    pr = Presentation(p)
    out = []
    for i, slide in enumerate(pr.slides):
        out.append("--- slide %d ---" % (i + 1))
        for sh in slide.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
    return "\n".join(out), len(pr.slides)


def extract_xlsx(p):
    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append("=== sheet: %s ===" % ws.title)
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                out.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(out), len(wb.worksheets)


def extract_xls(p):
    import xlrd
    wb = xlrd.open_workbook(p)
    out = []
    for ws in wb.sheets():
        out.append("=== sheet: %s ===" % ws.name)
        for r in range(ws.nrows):
            vals = [ws.cell_value(r, c) for c in range(ws.ncols)]
            if any(str(v).strip() for v in vals):
                out.append(" | ".join(str(v) for v in vals))
    return "\n".join(out), wb.nsheets


def extract_doc(p):
    """老 .doc：用 macOS 自带 textutil 转 txt。"""
    tmp = p + ".textutil.txt"
    try:
        subprocess.run(["textutil", "-convert", "txt", "-output", tmp, p], check=True,
                       capture_output=True, timeout=60)
        return read_text(tmp), None
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)


EXTRACTORS = {
    ".pdf": extract_pdf, ".docx": extract_docx, ".pptx": extract_pptx,
    ".xlsx": extract_xlsx, ".xls": extract_xls, ".doc": extract_doc,
}


def classify(ext):
    ext = ext.lower()
    if ext in EXTRACTORS:
        return "extract"
    if ext in TEXT_EXT:
        return "text"
    if ext in IMG_EXT:
        return "image"
    if ext in AV_EXT:
        return "av"
    if ext in ARC_EXT:
        return "archive"
    return "unsupported"


def process(src, manifest):
    ext = os.path.splitext(src)[1]
    kind = classify(ext)
    rel = os.path.relpath(src, os.path.expanduser("~"))
    dst = os.path.join(OUT_ROOT, rel + ".txt")
    if kind in ("image", "av", "archive", "unsupported"):
        manifest.append((rel, ext, kind, "", "", "agent-readable" if kind != "unsupported" else "needs-tool"))
        return
    if os.path.exists(dst) and os.path.getmtime(dst) >= os.path.getmtime(src):
        return
    os.makedirs(os.path.dirname(dst), exist_ok=True)
    pages = ""
    try:
        if kind == "extract":
            txt, pages = EXTRACTORS[ext](src)
        else:
            txt = read_text(src, 200000)
        with open(dst, "w", encoding="utf-8") as f:
            f.write("# 源：%s\n# 抽取：%s\n\n%s\n" % (rel, datetime.datetime.now().isoformat(timespec="seconds"), txt))
        status = "ok" if len(txt.strip()) > 30 else "empty?(可能扫描件，交由 agent 用 pdf/view_image 读)"
        manifest.append((rel, ext, kind, str(pages), str(len(txt)), status))
    except Exception as e:
        manifest.append((rel, ext, kind, "", "", "error: %s" % str(e)[:80]))


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    roots = sys.argv[1:]
    os.makedirs(OUT_ROOT, exist_ok=True)
    manifest = []
    for root in roots:
        root = os.path.abspath(os.path.expanduser(root))
        if os.path.isfile(root):
            process(root, manifest)
            continue
        for dp, _dn, fns in os.walk(root):
            for fn in fns:
                if fn.startswith("."):
                    continue
                process(os.path.join(dp, fn), manifest)
    mpath = os.path.join(OUT_ROOT, "_manifest.tsv")
    new = not os.path.exists(mpath)
    with open(mpath, "a", encoding="utf-8") as f:
        if new:
            f.write("file\text\tkind\tpages\tchars\tstatus\n")
        for row in manifest:
            f.write("\t".join(row) + "\n")
    ok = sum(1 for r in manifest if r[5] == "ok")
    print("处理 %d 个；成功抽文本 %d；其余标为 agent-readable / 待处理" % (len(manifest), ok))
    print("文本目录：%s" % OUT_ROOT)


if __name__ == "__main__":
    main()
